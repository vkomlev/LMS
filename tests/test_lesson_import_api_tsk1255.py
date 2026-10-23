"""tsk-1255: разбор файлов урока и публикация черновика через API.

Проверяем то, чего нет в разборщике: узлы встают под курс-родитель, лимит
попыток пре-теста ставится штатно, повторная публикация обновляет тот же
урок, а присланный черновик не может задеть чужие курсы.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
import pytest_asyncio
from docx import Document
from sqlalchemy import text

project_root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(project_root))

from app.core.config import Settings

pytestmark = pytest.mark.asyncio
_settings = Settings()


def _headers() -> dict[str, str]:
    return {"X-API-Key": next(iter(_settings.valid_api_keys))}


def _docx(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pretest(doc) -> None:
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "1", "Натрий — металл"
    t.cell(1, 0).text, t.cell(1, 1).text = "2", "Сера — металл"
    doc.add_paragraph("Ключ: 1) + 2) −")


def _homework(doc) -> None:
    doc.add_paragraph("К блоку №1")
    doc.add_paragraph("Выберите металлы:")
    for x in ("натрий", "сера", "медь"):
        doc.add_paragraph(x, style="List Paragraph")
    doc.add_paragraph("К блоку №2")
    doc.add_paragraph("Объясните, почему.")
    doc.add_paragraph("Ответы:")
    doc.add_paragraph("13", style="List Paragraph")


@pytest_asyncio.fixture(scope="function")
async def parent(db):
    """Курс-родитель с уникальным course_uid и uid сложности для подмены."""
    cid = (await db.execute(text(
        "INSERT INTO courses (title, access_level, course_uid) "
        "VALUES ('tsk1255 parent', 'manual_check', 'TSK1255-' || floor(random()*1e9)::text) RETURNING id"
    ))).scalar()
    uid = (await db.execute(text("SELECT course_uid FROM courses WHERE id = :i"), {"i": cid})).scalar()
    diff = (await db.execute(text("SELECT uid FROM difficulties ORDER BY id LIMIT 1"))).scalar()
    await db.commit()
    return {"id": cid, "uid": uid, "diff": diff}


async def _parse(client, course_id: int) -> dict:
    files = {
        "pretest": ("pre.docx", _docx(_pretest)),
        "homework": ("hw.docx", _docx(_homework)),
        "konspekt": ("kon.docx", _docx(lambda d: d.add_paragraph("БЛОК 1. Металлы"))),
    }
    r = await client.post(
        f"/api/v1/courses/{course_id}/lesson-imports/parse",
        data={"lesson_num": "3", "lesson_title": "Урок 3. Металлы",
              "video_url": "https://vkvideo.ru/video-1_2"},
        files=files, headers=_headers(),
    )
    assert r.status_code == 200, r.text
    return r.json()


async def test_parse_returns_draft_without_db_writes(client, db, parent):
    draft = await _parse(client, parent["id"])
    assert [n["suffix"] for n in draft["nodes"]] == ["PRE", "KON", "POST", "HW"]
    hw = {t["uid"]: t for t in draft["tasks"] if t["node"] == "HW"}
    assert hw["HW-01"]["task_content"]["type"] == "SA"
    assert hw["HW-02"]["task_content"]["type"] == "TA"
    n = (await db.execute(text("SELECT count(*) FROM courses WHERE course_uid LIKE :u"),
                          {"u": parent["uid"] + "-U%"})).scalar()
    assert n == 0


async def test_publish_creates_lesson_and_is_idempotent(client, db, parent):
    draft = await _parse(client, parent["id"])
    for t in draft["tasks"]:
        t["difficulty_uid"] = parent["diff"]
    url = f"/api/v1/courses/{parent['id']}/lesson-imports/publish"
    r = await client.post(url, json=draft, headers=_headers())
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["lesson_course_uid"] == f"{parent['uid']}-U03"
    assert res["tasks_created"] == 6 and res["materials_created"] == 2

    rows = (await db.execute(text(
        "SELECT external_uid, max_attempts, requirement_level, solution_rules->>'diagnostic' "
        "FROM tasks WHERE external_uid LIKE :u ORDER BY external_uid"
    ), {"u": parent["uid"] + "-U03-%"})).all()
    pre = [r for r in rows if "-PRE-" in r[0] or "-POST-" in r[0]]
    assert len(pre) == 4 and all(r[1] == 1 and r[2] == "recommended" and r[3] == "true" for r in pre)
    hw = [r for r in rows if "-HW-" in r[0]]
    assert len(hw) == 2 and all(r[1] is None for r in hw)

    parent_link = (await db.execute(text(
        "SELECT cp.parent_course_id FROM course_parents cp JOIN courses c ON c.id = cp.course_id "
        "WHERE c.course_uid = :u"), {"u": f"{parent['uid']}-U03"})).scalar()
    assert parent_link == parent["id"]

    draft["tasks"][0]["task_content"]["stem"] = "Натрий — щелочной металл"
    r2 = await client.post(url, json=draft, headers=_headers())
    assert r2.status_code == 200, r2.text
    assert r2.json()["tasks_created"] == 0 and r2.json()["tasks_updated"] == 6
    stem = (await db.execute(text(
        "SELECT task_content->>'stem' FROM tasks WHERE external_uid = :u"
    ), {"u": f"{parent['uid']}-U03-PRE-01"})).scalar()
    assert stem == "Натрий — щелочной металл"


async def test_publish_rejects_bad_draft_before_writing(client, db, parent):
    draft = await _parse(client, parent["id"])
    for t in draft["tasks"]:
        t["difficulty_uid"] = parent["diff"]
    draft["tasks"][0]["solution_rules"]["correct_options"] = ["нет-такого"]
    r = await client.post(f"/api/v1/courses/{parent['id']}/lesson-imports/publish",
                          json=draft, headers=_headers())
    assert r.status_code == 422, r.text
    n = (await db.execute(text("SELECT count(*) FROM courses WHERE course_uid LIKE :u"),
                          {"u": parent["uid"] + "-U%"})).scalar()
    assert n == 0


async def test_publish_rejects_foreign_uid(client, parent):
    draft = await _parse(client, parent["id"])
    draft["tasks"][0]["uid"] = "../OTHER"
    r = await client.post(f"/api/v1/courses/{parent['id']}/lesson-imports/publish",
                          json=draft, headers=_headers())
    assert r.status_code == 422


async def test_parse_requires_docx_and_role(client, parent):
    url = f"/api/v1/courses/{parent['id']}/lesson-imports/parse"
    r = await client.post(url, data={"lesson_num": "1", "lesson_title": "x"},
                          files={"konspekt": ("a.pdf", b"%PDF")}, headers=_headers())
    assert r.status_code == 422
    r = await client.post(url, data={"lesson_num": "1", "lesson_title": "x"},
                          files={"konspekt": ("a.docx", b"x")})
    assert r.status_code in (401, 403)
