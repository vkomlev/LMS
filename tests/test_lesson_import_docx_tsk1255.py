"""tsk-1255: девять особенностей авторских docx при импорте урока.

Каждая особенность — отдельный тест на маленьком синтетическом docx: так
видно, какая именно потеря вернулась. Список особенностей взят из README
курса «Дипломат», где ручные скрипты обкатаны на 12 уроках.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

project_root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(project_root))

from app.services.checking_service import CheckingService
from app.services.lesson_import import answers
from app.services.lesson_import.draft import build_draft, mixed_script_tokens
from app.services.lesson_import.homework import parse_homework
from app.services.lesson_import.konspekt import konspekt_html
from app.services.lesson_import.pretest import parse_pretest
from app.utils.exceptions import DomainError


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _hw_doc(*paras: tuple[str, str] | str):
    """Документ ДЗ из абзацев: строка — Normal, кортеж (стиль, текст)."""
    doc = Document()
    for p in paras:
        style, text = ("Normal", p) if isinstance(p, str) else p
        doc.add_paragraph(text, style=style)
    return doc


LP = "List Paragraph"


# 1. Над/подстрочные знаки → юникод.
def test_superscript_and_subscript_become_unicode():
    doc = Document()
    p = doc.add_paragraph("1s")
    p.add_run("2").font.superscript = True
    p.add_run("2s")
    p.add_run("2").font.superscript = True
    p.add_run(" H")
    p.add_run("2").font.subscript = True
    p.add_run("O")
    html = konspekt_html(_save(doc))
    assert "1s²2s² H₂O" in html


# 2. Символ служебного шрифта w:sym (стрелка реакции).
def test_w_sym_arrow_is_kept():
    doc = Document()
    p = doc.add_paragraph("A + B ")
    run = p.add_run()
    sym = OxmlElement("w:sym")
    sym.set(qn("w:font"), "Wingdings")
    sym.set(qn("w:char"), "F0E0")
    run._r.append(sym)
    p.add_run(" C")
    assert "A + B → C" in konspekt_html(_save(doc))


# 3. Стрелка обычным текстом в шрифте Wingdings.
def test_wingdings_text_arrow():
    doc = Document()
    p = doc.add_paragraph("Fe ")
    r = p.add_run("à")
    r.font.name = "Wingdings"
    p.add_run(" Fe²⁺")
    html = konspekt_html(_save(doc))
    assert "Fe → Fe²⁺" in html
    assert "à" not in html


# 4. Гиперссылки w:hyperlink: адрес не теряется.
def test_hyperlink_keeps_href():
    doc = Document()
    p = doc.add_paragraph("Рекомендуем ресурс: ")
    rid = doc.part.relate_to("https://xumuk.ru/esa/", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "xumuk"
    r.append(t)
    link.append(r)
    p._p.append(link)
    html = konspekt_html(_save(doc))
    assert '<a href="https://xumuk.ru/esa/"' in html and ">xumuk</a>" in html


# 5. Нумерация List Paragraph: номер в нумерации Word → варианты 1) 2).
def test_list_paragraph_options_numbered():
    doc = _hw_doc("К блоку №1", "Выберите металлы:", (LP, "натрий"), (LP, "– сера"), (LP, "медь"))
    b = parse_homework(_save(doc)).blocks[0]
    html = "".join(b.html)
    assert "<strong>1)</strong> натрий" in html
    assert "<strong>2)</strong> сера" in html  # тире автора снято
    assert "<strong>3)</strong> медь" in html


# 6. Соответствие строками с табуляцией; заголовок колонки склеивается до двоеточия.
def test_tabbed_matching_two_columns():
    doc = _hw_doc(
        "К блоку №2", "Установите соответствие:",
        "Элемент\tБлок в ПСХЭ",
        "А)\tРубидий\t1.\ts-элементы",
        "Б)\tХром\t2.\td-элементы",
        "В)\tМедь",
    )
    html = "".join(parse_homework(_save(doc)).blocks[0].html)
    assert "<p><strong>Элемент</strong></p>" in html
    assert "<strong>А)</strong> Рубидий" in html and "<strong>В)</strong> Медь" in html
    assert "<p><strong>Блок в ПСХЭ</strong></p>" in html
    assert "<strong>2)</strong> d-элементы" in html


def test_table_matching_joins_split_header():
    doc = _hw_doc("К блоку №1", "Установите соответствие:")
    t = doc.add_table(rows=4, cols=2)
    left = ["Формула", "вещества:", "А) NaCl", "Б) HCl"]
    right = ["Степень окисления", "хлора:", "1) −1", "2) +1"]
    for i in range(4):
        t.cell(i, 0).text = left[i]
        t.cell(i, 1).text = right[i]
    html = "".join(parse_homework(_save(doc)).blocks[0].html)
    assert "<strong>Формула вещества:</strong>" in html
    assert "<strong>Степень окисления хлора:</strong>" in html
    assert "<strong>А)</strong> NaCl" in html  # буква автора снята, своя поставлена
    assert "<strong>1)</strong> −1" in html
    assert "хлора:</li>" not in html


# 7. Однострочная таблица из ≥2 колонок — перечень, нумерация по колонкам.
def test_single_row_table_is_list_by_columns():
    doc = _hw_doc("К блоку №5", "Выберите оксиды:")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "CaO"
    t.cell(0, 0).add_paragraph("NaOH")
    t.cell(0, 1).text = "SO₂"
    html = "".join(parse_homework(_save(doc)).blocks[0].html)
    assert html.index("<strong>1)</strong> CaO") < html.index("<strong>2)</strong> NaOH")
    assert "<strong>3)</strong> SO₂" in html


def test_single_cell_table_is_plain_text():
    doc = _hw_doc("К блоку №5", "Условие:")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Длинный текст условия"
    html = "".join(parse_homework(_save(doc)).blocks[0].html)
    assert "<p>Длинный текст условия</p>" in html and "<ul>" not in html


# 8. Ключи «Ответы:»/«Ответ:», подпись к пустой таблице, номер блока из скобки.
def test_answer_keys_list_after_header():
    doc = _hw_doc(
        "К блоку №1", "Вопрос один", (LP, "а"), (LP, "б"),
        "(вместо задания к блоку №4)", "Вопрос два",
        "Ответы:", (LP, "13"), (LP, "Задание с развёрнутым ответом."), "Решение автора…",
    )
    res = parse_homework(_save(doc))
    assert [b.num for b in res.blocks] == [1, 4]
    assert [b.key for b in res.blocks] == ["13", None]
    assert "Решение" not in "".join(res.blocks[1].html)


def test_answer_label_over_empty_table_is_not_keys():
    doc = _hw_doc("К блоку №1", "Вопрос", "Ответ:")
    t = doc.add_table(rows=2, cols=3)
    for i, ch in enumerate("АБВ"):
        t.cell(0, i).text = ch
    doc.add_paragraph("К блоку №2")
    doc.add_paragraph("Второй вопрос")
    doc.add_paragraph("Ответ:")
    doc.add_paragraph("212", style=LP)
    doc.add_paragraph("35", style=LP)
    res = parse_homework(_save(doc))
    assert [b.num for b in res.blocks] == [1, 2]
    assert [b.key for b in res.blocks] == ["212", "35"]
    assert "Ответ" not in "".join(res.blocks[0].html)


def test_numbered_keys_in_header_paragraph():
    doc = _hw_doc("К блоку №1", "x", "К блоку №2", "y")
    p = doc.add_paragraph("Ответы:")
    p.add_run().add_break()
    p.add_run("1) 1423")
    doc.add_paragraph("2) 14")
    res = parse_homework(_save(doc))
    assert [b.key for b in res.blocks] == ["1423", "14"]


def test_key_count_mismatch_warns():
    doc = _hw_doc("К блоку №1", "x", "К блоку №2", "y", "Ответы:", (LP, "12"))
    res = parse_homework(_save(doc))
    assert any("ключей 1, блоков 2" in w for w in res.warnings)


# 9. Дробные ответы сверяются без strip_punctuation.
@pytest.mark.parametrize("key,wrong", [("67,2", "672"), ("0,10", "010")])
def test_decimal_key_rejects_lost_separator(key, wrong):
    rules = answers.short_answer_rules(key, 2)["short_answer"]
    steps = rules["normalization"]
    assert "strip_punctuation" not in steps
    accepted = [a["value"] for a in rules["accepted_answers"]]
    match = CheckingService._matches_short_answer
    assert any(match(key.replace(",", "."), a, steps) for a in accepted)
    assert not any(match(wrong, a, steps) for a in accepted)


def test_digit_key_accepts_commas_and_spaces():
    rules = answers.short_answer_rules("5312", 2)["short_answer"]
    accepted = [a["value"] for a in rules["accepted_answers"]]
    match = CheckingService._matches_short_answer
    for value in ("5312", "5,3,1,2", "5 3 1 2"):
        assert any(match(value, a, rules["normalization"]) for a in accepted), value


# Плюс сохраняемое: цвет фрагментов, кроме чёрного и «авто».
def test_color_kept_except_black():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("красный")
    rpr = r._r.get_or_add_rPr()
    for val in ("FF0000",):
        c = OxmlElement("w:color")
        c.set(qn("w:val"), val)
        rpr.append(c)
    r2 = p.add_run(" чёрный")
    c2 = OxmlElement("w:color")
    c2.set(qn("w:val"), "000000")
    r2._r.get_or_add_rPr().append(c2)
    html = konspekt_html(_save(doc))
    assert '<span style="color: #FF0000">красный</span>' in html
    assert "#000000" not in html


def test_html_in_author_text_is_escaped():
    doc = Document()
    doc.add_paragraph("если a < b и c > d")
    html = konspekt_html(_save(doc))
    assert "a &lt; b" in html and "<b" not in html.replace("<br", "")


def test_pretest_key_and_missing_key_warning():
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "1", "Натрий — металл"
    t.cell(1, 0).text, t.cell(1, 1).text = "2", "Сера — металл"
    doc.add_paragraph("Ключ: 1) + ")
    res = parse_pretest(_save(doc))
    assert [(s.num, s.correct) for s in res.statements] == [(1, True), (2, None)]
    assert any("нет ключа" in w for w in res.warnings)


def test_mixed_script_formula_detected():
    assert mixed_script_tokens("<p>КСl и Н₂О, но Сера и NaCl</p>") == ["КСl", "Н₂О"]


def test_broken_file_is_422():
    with pytest.raises(DomainError) as exc:
        konspekt_html(b"not a docx")
    assert exc.value.status_code == 422


# Сквозной: урок курса «Дипломат» против эталона ручных скриптов.
_MATERIALS = Path("D:/Work/CreateCourses/courses/himiya-diplomat")


@pytest.mark.skipif(not _MATERIALS.exists(), reason="нет исходников курса «Дипломат»")
def test_end_to_end_lesson_matches_reference():
    src = _MATERIALS / "materials" / "urok-06"
    pick = lambda part: next(p for p in src.glob("*.docx") if part in p.name).read_bytes()  # noqa: E731
    draft = build_draft(lesson_num=6, lesson_title="Урок 6", konspekt=pick("Конспект"),
                        pretest=pick("тест"), homework=pick("Домашнее"))
    ref = json.loads((_MATERIALS / "exports" / "himiya-diplomat-lms-import.json").read_text("utf-8"))
    ref_tasks = {t["external_uid"]: t for t in ref["tasks"] if "-U06-" in t["external_uid"]}
    mine = {f"HIM-DIPLOMAT-U06-{t['uid']}": t for t in draft["tasks"]}
    assert set(mine) == set(ref_tasks)
    for uid, t in mine.items():
        r = ref_tasks[uid]
        assert t["task_content"]["stem"] == r["task_content"]["stem"], uid
        assert t["task_content"]["type"] == r["task_content"]["type"], uid
        a, b = t["solution_rules"], r["solution_rules"]
        assert a.get("correct_options") == b.get("correct_options"), uid
        if "short_answer" in b:
            got = sorted(x["value"] for x in a["short_answer"]["accepted_answers"])
            assert got == sorted(x["value"] for x in b["short_answer"]["accepted_answers"]), uid
    kon = next(m for m in draft["materials"] if m["uid"] == "KONSPEKT")["content"]["text"]
    ref_kon = next(m for m in ref["materials"] if m["external_uid"] == "HIM-DIPLOMAT-U06-KONSPEKT")
    assert ref_kon["content"]["text"].endswith(kon)
    assert [n["suffix"] for n in draft["nodes"]] == ["PRE", "KON", "POST", "HW"]
