"""tsk-1255: чтение авторского docx без потерь — общая часть импорта урока.

Перенесено из ручных скриптов курса «Дипломат» (CreateCourses/.../exports,
build_import.py + build_lessons.py), обкатанных на 12 уроках. Класс дефектов,
который здесь закрыт: **всё, что в docx живёт не в обычном тексте `w:r`,
пропадает молча** — над/подстрочные знаки, символы служебных шрифтов (`w:sym`),
стрелка обычным текстом в шрифте Wingdings, гиперссылки (`w:hyperlink`).
"""

from __future__ import annotations

import io
import logging
from typing import Iterator

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.utils.exceptions import DomainError

logger = logging.getLogger(__name__)

SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")

#: Символы служебных шрифтов (Wingdings/Symbol) → юникод.
SYMBOLS: dict[str, str] = {
    "F0E0": "→",
    "F0AE": "→",
    "F0DF": "←",
    "F0AB": "↔",
    "F0B0": "°",
}

#: Маркер Word (именованный цвет) → HEX.
HIGHLIGHT_HEX: dict[str, str] = {
    "yellow": "#FFFF00", "green": "#00FF00", "cyan": "#00FFFF", "magenta": "#FF00FF",
    "blue": "#0000FF", "red": "#FF0000", "darkBlue": "#000080", "darkCyan": "#008080",
    "darkGreen": "#008000", "darkMagenta": "#800080", "darkRed": "#800000",
    "darkYellow": "#808000", "darkGray": "#808080", "lightGray": "#C0C0C0",
}

LIST_STYLE = "List Paragraph"


def open_docx(data: bytes, label: str) -> DocxDocument:
    """Открыть docx из байтов; битый файл — понятная ошибка 422, а не 500."""
    try:
        return Document(io.BytesIO(data))
    except Exception as exc:  # python-docx бросает разные исключения на мусор
        logger.warning("lesson_import: не открылся docx %s: %s", label, exc)
        raise DomainError(
            detail=f"Файл «{label}» не читается как документ Word (.docx)",
            status_code=422,
        ) from exc


def esc(text: str) -> str:
    """Экранировать HTML: голый «<» съедает остаток материала."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _sym_text(run_el) -> str:
    """Символы `w:sym` внутри фрагмента: обычный `run.text` их не отдаёт."""
    out: list[str] = []
    for sym in run_el.findall(qn("w:sym")):
        code = (sym.get(qn("w:char")) or "").upper()
        out.append(SYMBOLS.get(code, "→" if code.startswith("F0") else ""))
    return "".join(out)


def run_text(run_el) -> str:
    """Текст фрагмента `w:r` с индексами, символами и Wingdings-стрелкой."""
    # Табуляция и разрыв строки — отдельные элементы, не `w:t`: без них
    # соответствие, свёрстанное табуляцией, слипается в «ЭлементБлок».
    text = "".join(
        (el.text or "") if el.tag == qn("w:t") else ("\t" if el.tag == qn("w:tab") else "\n")
        for el in run_el
        if el.tag in (qn("w:t"), qn("w:tab"), qn("w:br"))
    )
    if not text.strip():
        if text:
            return text
        return _sym_text(run_el)
    rpr = run_el.find(qn("w:rPr"))
    if rpr is None:
        return text
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None and fonts.get(qn("w:ascii")) == "Wingdings":
        # Стрелка набрана обычным текстом в шрифте Wingdings: «à» (0xE0) — это F0E0.
        return "".join(SYMBOLS.get(f"F0{ord(c):02X}", "→") for c in text)
    va = rpr.find(qn("w:vertAlign"))
    val = va.get(qn("w:val")) if va is not None else None
    if val == "superscript":
        return text.translate(SUP)
    if val == "subscript":
        return text.translate(SUB)
    return text


def _runs(par: Paragraph) -> Iterator:
    """Все `w:r` абзаца, включая лежащие внутри гиперссылок."""
    for child in par._p:
        if child.tag == qn("w:r"):
            yield child
        elif child.tag == qn("w:hyperlink"):
            yield from child.findall(qn("w:r"))


def para_text(par: Paragraph) -> str:
    """Плоский текст абзаца (для разбора структуры, не для показа)."""
    return "".join(run_text(r) for r in _runs(par)).strip()


def colored(run_el, html: str) -> str:
    """Обернуть фрагмент в span с цветом автора; чёрный и «авто» не переносим."""
    if not html.strip():
        return html
    rpr = run_el.find(qn("w:rPr"))
    if rpr is None:
        return html
    style: list[str] = []
    color = rpr.find(qn("w:color"))
    val = color.get(qn("w:val")) if color is not None else None
    if val and val.lower() not in ("auto", "000000"):
        style.append(f"color: #{val}")
    hl = rpr.find(qn("w:highlight"))
    bg = HIGHLIGHT_HEX.get(hl.get(qn("w:val")) or "") if hl is not None else None
    if bg is None:
        shd = rpr.find(qn("w:shd"))
        fill = shd.get(qn("w:fill")) if shd is not None else None
        if fill and fill.lower() not in ("auto", "ffffff"):
            bg = f"#{fill}"
    if bg:
        style.append(f"background-color: {bg}")
    if not style:
        return html
    return f'<span style="{"; ".join(style)}">{html}</span>'


def hyperlink_targets(doc: DocxDocument) -> dict[str, str]:
    """rId → внешний адрес гиперссылки."""
    return {rid: rel.target_ref for rid, rel in doc.part.rels.items() if rel.is_external}


def para_html(par: Paragraph, targets: dict[str, str]) -> str:
    """Абзац как HTML: текст экранирован, цвет сохранён, ссылки — ссылками."""
    parts: list[str] = []
    for child in par._p:
        if child.tag == qn("w:r"):
            parts.append(colored(child, esc(run_text(child))))
        elif child.tag == qn("w:hyperlink"):
            href = targets.get(child.get(qn("r:id")) or "", "")
            text = "".join(run_text(r) for r in child.findall(qn("w:r")))
            if href:
                parts.append(
                    f'<a href="{esc(href)}" target="_blank" rel="noopener">{esc(text or href)}</a>'
                )
            else:
                parts.append(esc(text))
    return "".join(parts).strip().replace("\n", "<br>")


def style_name(par: Paragraph) -> str:
    """Имя стиля абзаца или пустая строка."""
    try:
        return par.style.name if par.style is not None else ""
    except (KeyError, ValueError):
        return ""


def blocks(doc: DocxDocument) -> list[Paragraph | Table]:
    """Абзацы и таблицы тела документа в порядке следования."""
    out: list[Paragraph | Table] = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            out.append(Paragraph(child, doc))
        elif isinstance(child, CT_Tbl):
            out.append(Table(child, doc))
    return out


def cell_text(cell) -> str:
    """Текст ячейки: абзацы через пробел."""
    return " ".join(t for t in (para_text(p) for p in cell.paragraphs) if t).strip()


def table_columns(table: Table) -> list[list[str]]:
    """Таблица → колонки абзацев (объединённые ячейки считаются один раз)."""
    cols: list[list[str]] = []
    for row in table.rows:
        seen: set[int] = set()
        for ci, cell in enumerate(row.cells):
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            paras = [t for t in (para_text(p) for p in cell.paragraphs) if t]
            while len(cols) <= ci:
                cols.append([])
            cols[ci].extend(paras)
    return [c for c in cols if c]
